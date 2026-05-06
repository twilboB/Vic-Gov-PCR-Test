"""
PCR POC — File scanner.
Converts campaign files to text and calls Gemini to classify + extract metrics.
"""

import json
import os
import re
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

import vertexai
from vertexai.generative_models import GenerativeModel

from config import GCP_PROJECT, GCP_LOCATION, GEMINI_MODEL, SCAN_PROMPT


def _init_vertex():
    vertexai.init(project=GCP_PROJECT, location=GCP_LOCATION)
    return GenerativeModel(GEMINI_MODEL)


# ── File → text converters ────────────────────────────────────────────────────

def _xlsx_to_text(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = "\t".join(cells).rstrip()
            if line:
                rows_text.append(line)
        if rows_text:
            parts.append(f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows_text))
    wb.close()
    return "\n\n".join(parts)


def _csv_to_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _pdf_to_text(path: Path) -> str:
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"=== Page {i} ===\n{text}")
    return "\n\n".join(pages)


def convert_file_to_text(path: Path) -> tuple[str, str | None]:
    """
    Returns (text_content, error_message).
    error_message is None on success.
    """
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xls"):
            return _xlsx_to_text(path), None
        elif suffix == ".csv":
            return _csv_to_text(path), None
        elif suffix == ".pdf":
            return _pdf_to_text(path), None
        else:
            return "", f"Unsupported file type: {suffix}"
    except Exception as exc:
        return "", str(exc)


# ── Gemini scan call ──────────────────────────────────────────────────────────

def _extract_json(raw: str) -> dict:
    """Strip markdown fences if present and parse JSON."""
    text = raw.strip()
    # Remove ```json ... ``` or ``` ... ```
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return json.loads(text)


def scan_file(path: Path, model: GenerativeModel | None = None) -> dict:
    """
    Scan a single file and return the Gemini classification result dict.
    Handles conversion, prompt construction, Gemini call, and JSON parsing.
    """
    if model is None:
        model = _init_vertex()

    file_text, error = convert_file_to_text(path)

    if error:
        return {
            "filename": path.name,
            "channel": "unknown",
            "vendor": None,
            "format_detected": "unknown",
            "status": "unreadable",
            "issues": [error],
            "metrics": {},
            "raw_summary": f"File could not be read: {error}",
        }

    # Truncate to ~100 k chars to stay within context limits
    MAX_CHARS = 100_000
    if len(file_text) > MAX_CHARS:
        file_text = file_text[:MAX_CHARS] + "\n\n[TRUNCATED — file exceeds 100 000 characters]"

    prompt = SCAN_PROMPT.format(file_text=file_text)

    try:
        response = model.generate_content(prompt)
        result = _extract_json(response.text)
        result["filename"] = path.name  # ensure filename is always from path
        return result
    except json.JSONDecodeError as exc:
        return {
            "filename": path.name,
            "channel": "unknown",
            "vendor": None,
            "format_detected": "unknown",
            "status": "unclassified",
            "issues": [f"JSON parse error from Gemini: {exc}", "Raw response saved in raw_summary"],
            "metrics": {},
            "raw_summary": response.text[:500] if "response" in dir() else "No response",
        }
    except Exception as exc:
        return {
            "filename": path.name,
            "channel": "unknown",
            "vendor": None,
            "format_detected": "unknown",
            "status": "unreadable",
            "issues": [f"Gemini API error: {exc}"],
            "metrics": {},
            "raw_summary": str(exc),
        }


# ── Folder scan ───────────────────────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".pdf"}


def scan_folder(folder_path: str, progress_callback=None) -> list[dict]:
    """
    Scan all supported files in folder_path.
    progress_callback(filename, index, total) is called before each file is scanned.
    Returns list of scan result dicts.
    """
    folder = Path(folder_path)
    if not folder.is_dir():
        raise ValueError(f"Not a directory: {folder_path}")

    files = sorted(
        [f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS]
    )

    # Also flag unsupported files
    unsupported = [
        f for f in folder.iterdir()
        if f.is_file() and f.suffix.lower() not in SUPPORTED_EXTENSIONS
        and not f.name.startswith(".")
    ]

    model = _init_vertex()
    results = []

    for i, file_path in enumerate(files):
        if progress_callback:
            progress_callback(file_path.name, i, len(files))
        result = scan_file(file_path, model=model)
        results.append(result)

    for f in unsupported:
        results.append({
            "filename": f.name,
            "channel": "unknown",
            "vendor": None,
            "format_detected": "unknown",
            "status": "unreadable",
            "issues": [f"Unsupported file type: {f.suffix}"],
            "metrics": {},
            "raw_summary": "File type not supported by scanner.",
        })

    return results


# ── Quick CLI test ────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python scanner.py <path_to_file_or_folder>")
        sys.exit(1)

    target = Path(sys.argv[1])
    if target.is_dir():
        results = scan_folder(str(target), progress_callback=lambda name, i, n: print(f"  [{i+1}/{n}] {name}"))
        for r in results:
            print(json.dumps(r, indent=2))
    else:
        model = _init_vertex()
        result = scan_file(target, model=model)
        print(json.dumps(result, indent=2))
